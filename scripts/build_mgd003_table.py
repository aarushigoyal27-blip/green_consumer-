"""Extract the MCS MGD 003 'PV Only' self-consumption shares into a tidy CSV.

Source: config/raw/MGD003-LookupTables-FINAL.xlsx (MCS MGD 003 look-up tables,
https://mcscertified.com/?standard=look-up-tables-pv-self-consumption).
Output: config/mcs_mgd003_pv_only.csv with one row per
(archetype, annual consumption band, annual generation band).
Usage: python scripts/build_mgd003_table.py [--check]
"""
import csv
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "config/raw/MGD003-LookupTables-FINAL.xlsx"
OUT = ROOT / "config/mcs_mgd003_pv_only.csv"
HEADER = re.compile(r"Occupancy: (.*?)\. Annual electricity consumption: ([\d,]+) kWh to ([\d,]+) kWh")
COLUMNS = ["archetype", "consumption_kwh_min", "consumption_kwh_max",
           "generation_kwh_min", "generation_kwh_max", "self_consumption_share"]


def _int(s):
    return int(s.replace(",", ""))


def build():
    ws = load_workbook(RAW, read_only=True, data_only=True)["Lookup Tables"]
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for i, r in enumerate(rows):
        for j, c in enumerate(r):
            m = HEADER.match(c) if isinstance(c, str) else None
            if not m:
                continue
            if (rows[i + 2][j], rows[i + 3][j]) != ("PV", "Only"):
                sys.exit(f"unexpected layout at row {i + 1}, col {j + 1}")
            for k in range(i + 4, len(rows)):
                label, share = rows[k][j - 1], rows[k][j]
                if not isinstance(label, str) or "kWh" not in label:
                    break
                g = re.findall(r"[\d,]+", label)
                out.append([m.group(1), _int(m.group(2)), _int(m.group(3)),
                            _int(g[0]), _int(g[1]), round(float(share), 6)])
    if len(out) != 540:
        sys.exit(f"expected 540 rows (3 archetypes x 9 demand bands x 20 generation bands), got {len(out)}")
    return out


if __name__ == "__main__":
    rows = build()
    if "--check" in sys.argv:
        existing = list(csv.reader(OUT.open(newline="", encoding="utf-8")))[1:]
        if [list(map(str, r)) for r in rows] != existing:
            sys.exit("config/mcs_mgd003_pv_only.csv differs from the MGD 003 workbook")
        print("config/mcs_mgd003_pv_only.csv matches the MGD 003 workbook.")
    else:
        with OUT.open("w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(COLUMNS)
            w.writerows(rows)
        print(f"wrote {OUT.relative_to(ROOT)} ({len(rows)} rows)")
